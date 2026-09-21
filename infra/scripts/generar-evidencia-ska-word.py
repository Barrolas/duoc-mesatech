"""Genera Word de evidencia Skarlet (React, negocio, pruebas §11 #5).

Uso:
  python generar-evidencia-ska-word.py
  python generar-evidencia-ska-word.py "C:\\Users\\...\\Pictures\\Screenshots"

Por defecto lee PNG con prefijo SKA en Pictures\\Screenshots (equipo EV1).
"""

import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

from evidencia_docx import REGISTRY, SKA_FIGURES, add_catalog_figure
from mesatech_informe_estilo import add_indice, add_portada, setup_document

DEFAULT_BASE = Path(r"C:\Users\Administrador\Pictures\Screenshots")
BASE = Path(sys.argv[1]).resolve() if len(sys.argv) > 1 else DEFAULT_BASE
OUT_NAME = "MesaTech_EV1_Evidencia_Skarlet_Negocio_UI_Pruebas.docx"
OUT = BASE / OUT_NAME
OUT_FALLBACK = BASE / "MesaTech_EV1_Evidencia_Skarlet_Negocio_UI_Pruebas_v2.docx"

FIGURES = [f for f in SKA_FIGURES if f.get("cap") in (6, 7)]


def main():
    BASE.mkdir(parents=True, exist_ok=True)
    doc = Document()
    setup_document(doc)
    add_portada(doc)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    from mesatech_informe_estilo import add_run_styled, COLOR_ACENTO

    add_run_styled(
        p,
        "Anexo — Evidencia React / pruebas SKA (mismo estilo que informe principal)",
        bold=True,
        color=COLOR_ACENTO,
    )
    doc.add_paragraph()

    REGISTRY.reset()
    for item in FIGURES:
        add_catalog_figure(doc, {**item, "cap": 12}, BASE, registry=REGISTRY)

    try:
        doc.save(OUT)
        print("Generado:", OUT)
    except PermissionError:
        doc.save(OUT_FALLBACK)
        print("Generado (cierra el .docx anterior si estaba abierto):", OUT_FALLBACK)

    missing = []
    for item in FIGURES:
        name = item.get("file") or (item.get("files") or [""])[0]
        if name and not (BASE / name).is_file() and name not in missing:
            missing.append(name)
    if missing:
        print("Capturas obligatorias faltantes:", len(missing))
        for name in missing:
            print(" -", name)
    else:
        print("Todas las capturas obligatorias SKA están incluidas.")


if __name__ == "__main__":
    main()
