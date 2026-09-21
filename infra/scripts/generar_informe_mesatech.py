"""Genera informe EP1 MesaTech Cloud: narrativa del caso + evidencias numeradas MT."""

from __future__ import annotations

import sys
from pathlib import Path

from docx import Document

from evidencia_docx import (
    REGISTRY,
    FigureRegistry,
    NIN_FIGURES,
    NIC_FIGURES,
    SKA_FIGURES,
    add_indice_evidencias,
    append_chapter_figures,
    append_entra_from_source,
)
from mesatech_informe_contenido import (
    anexo_local,
    cap1_contexto,
    cap1_dominio,
    cap1_escenarios_uso,
    cap1_propuesta,
    cap10_e2e_pasos,
    cap11_conclusiones,
    cap2_entra_narrativa,
    cap3_gateway_narrativa,
    cap4_infra_narrativa,
    cap5_backend_narrativa,
    cap6_react_narrativa,
    cap7_autorizacion_narrativa,
    cap8_versionamiento_narrativa,
    cap9_marca_narrativa,
    matriz_funcional_ep1,
    mapeo_componente_negocio,
)
from mesatech_informe_estilo import (
    COLOR_PRIMARIO,
    add_body_paragraphs,
    add_branded_table,
    add_chapter_heading,
    add_figure_block,
    add_indice,
    add_portada,
    add_run_styled,
    add_subsection_heading,
    setup_document,
)

REPO_ROOT = Path(__file__).resolve().parents[2]
DEFAULT_SRC = REPO_ROOT / "Screenshots EP1.docx"
DEFAULT_ASSETS = Path(r"C:\Users\Administrador\Pictures\Screenshots")
OUT_NAME = "MesaTech_EV1_Informe.docx"
DIAGRAMA_DESPLIEGUE = REPO_ROOT / "docs/assets/diagramas/arquitectura_mesatech_ev1_despliegue.png"
DIAGRAMA_PROPUESTA = REPO_ROOT / "docs/assets/diagramas/arquitectura_mesatech_propuesta.png"


def add_cap1_caso(doc: Document, reg: FigureRegistry) -> None:
    add_chapter_heading(
        doc,
        "1. Contexto del caso y propuesta de solución",
        "DSY1107 EP1 — MesaTech Cloud. Informe pormenorizado de la solución implementada "
        "por el equipo, alineado a la guía oficial y a la rúbrica §12–§15.",
    )
    ctx = cap1_contexto()
    add_subsection_heading(doc, "1.1 Contexto de negocio")
    add_body_paragraphs(doc, ctx["contexto"])
    add_subsection_heading(doc, "1.2 Desafío y objetivos de la evaluación")
    add_body_paragraphs(doc, ctx["desafio"])
    add_subsection_heading(doc, "1.3 Actores y responsabilidades")
    add_body_paragraphs(doc, ctx["actores"])
    add_subsection_heading(doc, "1.4 Alcance funcional (guía EP1 §4.1)")
    add_body_paragraphs(
        doc,
        [
            "La guía EP1 define qué puede hacer cada perfil sobre las solicitudes de soporte "
            "MesaTech. La implementación respeta esta matriz en el BFF (403) y en la "
            "navegación React.",
        ],
    )
    add_branded_table(
        doc,
        ("Acción", "Cliente", "Operador", "Administrador"),
        matriz_funcional_ep1(),
    )
    add_subsection_heading(doc, "1.5 Escenarios de uso MesaTech Cloud")
    add_branded_table(
        doc,
        ("Actor", "Situación de soporte", "Resultado en la plataforma"),
        cap1_escenarios_uso(),
    )
    add_subsection_heading(doc, "1.6 Propuesta técnica y mapeo al negocio")
    add_body_paragraphs(doc, cap1_propuesta())
    add_branded_table(
        doc,
        ("Componente", "Rol en MesaTech", "Aporte al servicio de soporte"),
        mapeo_componente_negocio(),
    )
    add_subsection_heading(doc, "1.7 Dominio: solicitudes, estados y reglas")
    add_body_paragraphs(doc, cap1_dominio())

    fig_id, code = reg.allocate(
        1,
        "Diagrama de despliegue EV1 (EC2, Gateway, flujos)",
        "§12 diagrama",
        key="diagrama_despliegue",
    )
    diagram_path = DIAGRAMA_DESPLIEGUE if DIAGRAMA_DESPLIEGUE.is_file() else DIAGRAMA_PROPUESTA
    add_figure_block(
        doc,
        fig_id,
        "Arquitectura implementada MesaTech Cloud",
        "Componentes en AWS y Microsoft Entra ID; React solo consume API Gateway. "
        "BFF en EC2; microservicios y PostgreSQL en capas separadas.",
        "§12 / §14 Arquitectura",
        image_paths=[diagram_path] if diagram_path.is_file() else [],
        evidencia_code=code,
    )


def add_cap10_cap11(doc: Document, reg: FigureRegistry) -> None:
    add_chapter_heading(
        doc,
        "10. Flujo extremo a extremo, pruebas §11 y checklist §15",
        "Demostración del recorrido §10 de la guía EP1, resultados de las pruebas mínimas "
        "y cumplimiento del checklist de entrega.",
    )

    add_subsection_heading(doc, "10.1 Flujo extremo a extremo (§10 guía EP1)")
    add_body_paragraphs(
        doc,
        [
            "El flujo se interpreta sobre un caso real de MesaTech: un cliente registra un "
            "incidente, un operador lo toma y avanza estados hasta la resolución, y la "
            "información persiste para consultas posteriores. La tabla condensa los doce "
            "pasos oficiales; las figuras SKA/NIN/NIC del informe ilustran cada capa.",
        ],
    )
    add_branded_table(
        doc,
        ("Paso", "Descripción"),
        cap10_e2e_pasos(),
    )

    add_subsection_heading(doc, "10.2 Pruebas mínimas obligatorias (§11)")
    add_indice_evidencias(doc, reg)
    add_branded_table(
        doc,
        ("#", "Prueba §11", "Resultado / evidencia"),
        (
            ("1", "Sin autenticación", f"UI sin datos privados · {reg.ref('sin_autenticacion')}"),
            (
                "2",
                "Gateway sin token",
                f"401 Unauthorized · {reg.ref('prueba_gw_sin_jwt')}",
            ),
            (
                "3",
                "JWT inválido",
                f"401 · {reg.ref('prueba_jwt_invalido')} / {reg.ref('prueba_jwt_invalido_ska')}",
            ),
            (
                "4",
                "JWT válido + operación autorizada",
                f"200 · {reg.ref('prueba_jwt_valido')}",
            ),
            (
                "5",
                "Acción no permitida (rol)",
                f"403 · {reg.ref('auth_403_ui')} / {reg.ref('auth_403_api')}",
            ),
            ("6", "CORS desde React", f"Sin bloqueo CORS · {reg.ref('prueba_cors_navegador')}"),
            (
                "7",
                "Regla de estado inválida",
                f"409 Conflict · {reg.ref('regla_estado_resuelta')} / {reg.ref('regla_estado_cerrada')}",
            ),
            (
                "8",
                "Coexistencia v1 y v2",
                f"{reg.ref('versionamiento_v1')} / {reg.ref('versionamiento_v2')}",
            ),
            (
                "9",
                "Persistencia",
                f"Datos en UI y reconsulta API · {reg.ref('crear_solicitud')} / "
                f"{reg.ref('listar_mias')} / {reg.ref('persistencia_postman')} / "
                f"{reg.ref('infra_postgres')}",
            ),
        ),
    )

    add_subsection_heading(doc, "10.3 Checklist antes de entregar (§15)")
    add_branded_table(
        doc,
        ("Requisito §15", "Cómo se cumple / evidencia"),
        (
            ("Login / logout Entra", f"MSAL en React · Cap. 2 y {reg.ref('sin_autenticacion')}"),
            ("Access token para API", reg.ref("prueba_jwt_valido")),
            ("Front solo vía Gateway", reg.ref("gw_invoke_url")),
            ("CORS operativo", f"{reg.ref('gw_cors')} / {reg.ref('prueba_cors_navegador')}"),
            ("Gateway rechaza sin JWT", reg.ref("prueba_gw_sin_jwt")),
            ("BFF revalida JWT", reg.ref("bff_401_sin_jwt")),
            ("BFF sin BD; proxy MS", reg.ref("infra_bff")),
            (
                "Dos MS con persistencia",
                f"{reg.ref('infra_ms_solicitudes')} / {reg.ref('infra_ms_catalogo')} / "
                f"{reg.ref('persistencia_postman')}",
            ),
            ("Rutas v1 + v2", "Cap. 8"),
            ("Despliegue EC2", "Cap. 4 y 5"),
            ("Diferencias por rol", reg.ref("auth_403_api")),
            ("Sin secretos en Git", "Contraseñas y tokens redactados en informe y repo"),
        ),
    )

    add_subsection_heading(doc, "10.4 Cobertura rúbrica §14 (resumen)")
    add_branded_table(
        doc,
        ("Área §14", "Evidencia en este informe"),
        (
            ("React + MSAL", "Cap. 6 — login, token, consumo Gateway"),
            ("Microsoft Entra ID", "Cap. 2 — registros SPA/API, roles, scope"),
            ("JWT y claims", "Cap. 2 y 6 — /api/usuario y Network"),
            ("AWS API Gateway", "Cap. 3 — rutas, JWT Authorizer, CORS"),
            ("Spring Security", "Cap. 5 — Resource Server BFF, 401/403"),
            ("Arquitectura", f"Cap. 1 {reg.ref('diagrama_despliegue', '')} y cap. 4"),
            ("Integración y negocio", "Cap. 6–7, pruebas §11, reglas de estado"),
        ),
    )

    add_chapter_heading(doc, "11. Conclusiones", None)
    add_body_paragraphs(doc, cap11_conclusiones())

    add_chapter_heading(doc, "Anexo A — Entorno local y buenas prácticas", None)
    add_body_paragraphs(doc, anexo_local())


def build_informe(
    source_ep1: Path,
    asset_base: Path,
    out_path: Path,
) -> Document:
    REGISTRY.reset()
    source = Document(str(source_ep1))
    doc = Document()
    setup_document(doc)
    add_portada(doc)
    add_indice(doc)

    reg = REGISTRY
    add_cap1_caso(doc, reg)

    add_chapter_heading(
        doc,
        "2. Microsoft Entra ID e identidad del servicio",
        "Autenticación corporativa y emisión de tokens para operar MesaTech Cloud con "
        "roles alineados a cliente, operador y administrador de soporte.",
    )
    add_body_paragraphs(doc, cap2_entra_narrativa())
    append_entra_from_source(doc, source, registry=reg, skip_chapter_heading=True)

    add_chapter_heading(
        doc,
        "3. API Gateway (AWS HTTP API)",
        cap3_gateway_narrativa()[0],
    )
    add_body_paragraphs(doc, cap3_gateway_narrativa()[1:])
    append_chapter_figures(doc, asset_base, 3, "", "", NIN_FIGURES, registry=reg)

    add_chapter_heading(
        doc,
        "4. Infraestructura EC2 y PostgreSQL",
        cap4_infra_narrativa()[0],
    )
    add_body_paragraphs(doc, cap4_infra_narrativa()[1:])
    append_chapter_figures(doc, asset_base, 4, "", "", NIC_FIGURES, registry=reg)

    add_chapter_heading(
        doc,
        "5. Backend BFF y microservicios",
        cap5_backend_narrativa()[0],
    )
    add_body_paragraphs(doc, cap5_backend_narrativa()[1:])
    append_chapter_figures(doc, asset_base, 5, "", "", NIC_FIGURES, registry=reg)

    add_chapter_heading(
        doc,
        "6. Aplicación React y MSAL",
        cap6_react_narrativa()[0],
    )
    add_body_paragraphs(doc, cap6_react_narrativa()[1:])
    append_chapter_figures(doc, asset_base, 6, "", "", SKA_FIGURES, registry=reg)

    add_chapter_heading(
        doc,
        "7. Autorización, roles y pruebas de API",
        cap7_autorizacion_narrativa()[0],
    )
    add_body_paragraphs(doc, cap7_autorizacion_narrativa()[1:])
    append_chapter_figures(doc, asset_base, 7, "", "", SKA_FIGURES, registry=reg)

    add_chapter_heading(
        doc,
        "8. Versionamiento v1 / v2",
        cap8_versionamiento_narrativa()[0],
    )
    add_body_paragraphs(doc, cap8_versionamiento_narrativa()[1:])
    append_chapter_figures(doc, asset_base, 8, "", "", NIN_FIGURES, registry=reg)

    add_chapter_heading(
        doc,
        "9. Identidad corporativa",
        cap9_marca_narrativa()[0],
    )
    add_body_paragraphs(doc, cap9_marca_narrativa()[1:])
    append_chapter_figures(doc, asset_base, 9, "", "", NIN_FIGURES, registry=reg)

    add_cap10_cap11(doc, reg)
    doc.save(out_path)
    return doc


def main():
    source = Path(sys.argv[1]).resolve() if len(sys.argv) > 1 else DEFAULT_SRC
    assets = Path(sys.argv[2]).resolve() if len(sys.argv) > 2 else DEFAULT_ASSETS
    out = Path(sys.argv[3]).resolve() if len(sys.argv) > 3 else REPO_ROOT / OUT_NAME

    if not source.is_file():
        print("No se encontró fuente Entra/embed:", source)
        sys.exit(1)

    try:
        doc = build_informe(source, assets, out)
    except PermissionError:
        out = out.with_name(out.stem + "_v2.docx")
        doc = build_informe(source, assets, out)
        print("(Archivo anterior abierto; guardado como _v2)")

    imgs = sum(1 for r in doc.part.rels.values() if "image" in r.reltype)
    print("Generado:", out)
    print("Imágenes totales:", imgs)
    print("Evidencias registradas:", len(REGISTRY.entries))


if __name__ == "__main__":
    main()
