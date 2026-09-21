"""Genera Word de evidencia Ninna (Gateway, CORS, JWT, v1/v2, marca).

Uso:
  python generar-evidencia-ninna-word.py
  python generar-evidencia-ninna-word.py "C:\\Users\\...\\Pictures\\Screenshots"

Por defecto lee PNG con prefijo NIN en Pictures\\Screenshots (nombres del equipo EV1).
"""

import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

from evidencia_docx import REGISTRY, NIN_FIGURES, add_catalog_figure
from mesatech_informe_estilo import add_portada, setup_document

DEFAULT_BASE = Path(r"C:\Users\Administrador\Pictures\Screenshots")
BASE = Path(sys.argv[1]).resolve() if len(sys.argv) > 1 else DEFAULT_BASE
OUT = BASE / "MesaTech_EV1_Evidencia_Ninna_Gateway_Marca.docx"

FIGURES_LEGACY = [
    {
        "file": "NIN-01 API Created.png",
        "id": "Figura NIN-01",
        "titulo": "HTTP API Gateway — API creada e Invoke URL",
        "texto": (
            "Consola AWS API Gateway (región us-east-1). Se crea la HTTP API del "
            "proyecto MesaTech EV1. Debe verse el nombre de la API, el API ID y la "
            "Invoke URL (https://….execute-api.us-east-1.amazonaws.com). Esa URL es "
            "la base de REACT_APP_API_BASE_URL en entrega: el React nunca apunta a "
            "la IP del BFF ni de los microservicios. EP1 §12 — API Gateway."
        ),
        "rubrica": "§12 HTTP API",
    },
    {
        "files": ["NIN-02 Rutas A.png", "NIN-02 Rutas B.png"],
        "id": "Figura NIN-02",
        "titulo": "Rutas expuestas en el Gateway",
        "texto": (
            "Listado de rutas (Routes) del HTTP API: paths de negocio alineados al "
            "BFF — como mínimo /api/usuario, /v1/solicitudes/mias, /v1/catalogo y "
            "/v2/solicitudes/mias. Los métodos HTTP deben coincidir con los que usa "
            "Axios en frontend/src/api/http.js. Referencia: "
            "docs/equipo/rutas-api-gateway-bff.md."
        ),
        "rubrica": "§12 Rutas",
    },
    {
        "files": ["NIN-03 Integrations A.png", "NIN-03 Integrations B.png"],
        "id": "Figura NIN-03",
        "titulo": "Integración HTTP hacia el BFF en EC2",
        "texto": (
            "Integración tipo HTTP URI hacia el Backend For Frontend desplegado en "
            "EC2 mesatech-ev1-bff (:8080, EIP 52.203.138.49). El Gateway reenvía el "
            "path sin alterar la semántica esperada por Spring Boot. Tras la petición "
            "autorizada, el BFF valida de nuevo el JWT (Resource Server) y proxifica "
            "a microservicios en red privada. §12 integraciones."
        ),
        "rubrica": "§12 Integración BFF",
    },
    {
        "file": "NIN-04 JWT.png",
        "id": "Figura NIN-04",
        "titulo": "JWT Authorizer (Microsoft Entra ID)",
        "texto": (
            "Authorizer JWT configurado con Issuer "
            "(https://login.microsoftonline.com/<tenant>/v2.0) y Audience (Client ID "
            "de la aplicación API en Entra), identity source Authorization. Rutas de "
            "negocio asociadas al authorizer. Sin token válido el Gateway responde "
            "401 y no ejecuta la integración al BFF. §12 / §14."
        ),
        "rubrica": "§12 JWT Authorizer",
    },
    {
        "file": "NIN-05 CORS.png",
        "id": "Figura NIN-05",
        "titulo": "CORS en HTTP API",
        "texto": (
            "Configuración CORS: Allow origins (http://localhost:3000 para CRA en "
            "desarrollo), headers Authorization y Content-Type, métodos GET, POST, "
            "PATCH, PUT, DELETE, OPTIONS. Permite que el navegador consuma la API "
            "desde React sin bloqueo cross-origin. §12 CORS."
        ),
        "rubrica": "§12 CORS Gateway",
    },
    {
        "file": "NIN-06-07 401 Unauthorized.png",
        "id": "Figura NIN-06",
        "titulo": "Gateway rechaza petición sin JWT (401)",
        "texto": (
            "Prueba GET /v1/solicitudes/mias contra la Invoke URL sin header "
            "Authorization (Postman o curl). Respuesta HTTP 401 Unauthorized. "
            "Evidencia §15 — el entrypoint público exige autenticación antes del "
            "BFF. Complementa NIC-07 (401 en BFF directo)."
        ),
        "rubrica": "§15 Gateway sin JWT",
    },
    {
        "file": "NIN-06-07 401 Unauthorized.png",
        "id": "Figura NIN-07",
        "titulo": "Gateway rechaza JWT inválido (401)",
        "texto": (
            "Misma ruta con Authorization: Bearer invalid (o token corrupto). "
            "El authorizer JWT del Gateway no valida firma/claims → 401. "
            "Prueba §11 #3. (Captura combinada NIN-06-07 en una sola imagen.)"
        ),
        "rubrica": "§15 JWT inválido",
    },
    {
        "file": "NIN-08-09 Browser Network.png",
        "id": "Figura NIN-08",
        "titulo": "Acceso autorizado vía Gateway (200)",
        "texto": (
            "Pestaña Network del navegador: XHR a execute-api con status 200 "
            "(usuario, catalogo, mias, solicitudes). Evidencia de consumo protegido "
            "con JWT MSAL. §11 #4 / §15."
        ),
        "rubrica": "§15 JWT válido",
    },
    {
        "file": "NIN-08-09 Browser Network.png",
        "id": "Figura NIN-09",
        "titulo": "React consumiendo Gateway sin error CORS",
        "texto": (
            "frontend/.env con REACT_APP_API_BASE_URL = Invoke URL. Tras login "
            "MSAL, preflight 204 y XHR 200 al dominio execute-api, sin blocked by "
            "CORS. §11 #6. (Misma captura Network que NIN-08.)"
        ),
        "rubrica": "§15 CORS navegador",
    },
    {
        "file": "NIN-10 v1.png",
        "id": "Figura NIN-10",
        "titulo": "Versionamiento — respuesta v1",
        "texto": (
            "GET /v1/solicitudes/mias vía Gateway con el mismo Bearer. Cuerpo: "
            "array JSON de solicitudes (contrato v1). §15 coexistencia v1."
        ),
        "rubrica": "§15 v1",
    },
    {
        "file": "NIN-11 v2.png",
        "id": "Figura NIN-11",
        "titulo": "Versionamiento — respuesta v2",
        "texto": (
            "GET /v2/solicitudes/mias con el mismo token. Cuerpo: objeto con "
            "version, usuario y solicitudes (contrato v2 del BFF). Contraste con "
            "NIN-10 demuestra versionamiento sin retirar v1. §14 / §11 #8."
        ),
        "rubrica": "§14 versionamiento",
    },
    {
        "file": "NIN-CRUD.png",
        "id": "Figura NIN-CRUD",
        "titulo": "Catálogo — CRUD administrador (MVP EP1)",
        "texto": (
            "Rol administrador en React: alta, edición y eliminación de categorías "
            "y prioridades vía Gateway (POST/PUT/DELETE /v1/catalogo/...). "
            "BFF aplica 403 si el token no es admin; ms-catalogo persiste en PostgreSQL. "
            "Responsable despliegue y rutas Gateway: Nico. §4.1 gestión catálogo."
        ),
        "rubrica": "§4.1 Catálogo admin",
    },
    {
        "file": "NIN-H1 Marca paleta logo.png",
        "id": "Figura NIN-H1",
        "titulo": "Identidad corporativa MesaTech Cloud",
        "texto": (
            "Paleta de colores, logo e isotipo definidos para el caso MesaTech "
            "(bloque H). Guía en docs/assets/marca/guia-marca.md; aplicación "
            "en React (App.css, logo.svg) e informe/presentación del equipo."
        ),
        "rubrica": "Informe — marca",
    },
]


def _paths_for(item):
    if "files" in item:
        return [BASE / name for name in item["files"]]
    return [BASE / item["file"]]


def _expected_names(item):
    if "files" in item:
        return item["files"]
    return [item["file"]]


def add_figure(doc, item, width_cm=15.5):
    paths = _paths_for(item)
    doc.add_heading(item["id"], level=2)
    p = doc.add_paragraph()
    r = p.add_run(item["titulo"])
    r.bold = True
    doc.add_paragraph(item["texto"])

    inserted = False
    for path in paths:
        if not path.is_file():
            continue
        inserted = True
        if len(paths) > 1:
            sub = doc.add_paragraph()
            sub.add_run(path.name).italic = True
        para = doc.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para.add_run().add_picture(str(path), width=Cm(width_cm))

    if inserted:
        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cr = cap.add_run(f"{item['id']} — {item['rubrica']}")
        cr.italic = True
        cr.font.size = Pt(9)
    else:
        names = ", ".join(f"«{n}»" for n in _expected_names(item))
        doc.add_paragraph(f"[Pendiente: {names} en {BASE}]")
    doc.add_paragraph()


def main():
    BASE.mkdir(parents=True, exist_ok=True)
    doc = Document()
    setup_document(doc)
    add_portada(doc)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    from mesatech_informe_estilo import COLOR_ACENTO, add_run_styled

    add_run_styled(
        p,
        "Anexo — Gateway AWS, CORS, JWT y versionamiento (estilo informe MesaTech)",
        bold=True,
        color=COLOR_ACENTO,
    )
    doc.add_paragraph()

    REGISTRY.reset()
    for item in NIN_FIGURES:
        add_catalog_figure(doc, {**item, "cap": 12}, BASE, registry=REGISTRY)

    doc.save(OUT)
    print("Generado:", OUT)

    missing = []
    for item in NIN_FIGURES:
        for name in (item.get("files") or [item["file"]]):
            if not (BASE / name).is_file() and name not in missing:
                missing.append(name)
    if missing:
        print("Capturas pendientes:", len(missing))
        for name in missing:
            print(" -", name)


if __name__ == "__main__":
    main()
