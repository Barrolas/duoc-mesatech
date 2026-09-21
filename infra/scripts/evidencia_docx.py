"""Catálogo de evidencias EP1 — numeración por capítulo (sin prefijos por integrante)."""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

from mesatech_informe_estilo import (
    COLOR_PRIMARIO,
    add_branded_table,
    add_chapter_heading,
    add_figure_block,
    add_run_styled,
)

# Identificador neutral: MT-{capítulo}-{orden}, p. ej. MT-3-06 (no ARI/NIN/NIC/SKA).


@dataclass
class FigureRegistry:
    """Registro de figuras del informe (numeración capítulo.orden + código MT)."""

    _order_by_cap: dict[int, int] = field(default_factory=dict)
    by_key: dict[str, str] = field(default_factory=dict)
    by_key_code: dict[str, str] = field(default_factory=dict)
    entries: list[tuple[str, str, str, str]] = field(default_factory=list)

    def reset(self) -> None:
        self._order_by_cap.clear()
        self.by_key.clear()
        self.by_key_code.clear()
        self.entries.clear()

    def allocate(self, cap: int, titulo: str, rubrica: str, key: str | None = None) -> tuple[str, str]:
        n = self._order_by_cap.get(cap, 0) + 1
        self._order_by_cap[cap] = n
        fig_id = f"Figura {cap}.{n}"
        code = f"MT-{cap}-{n:02d}"
        self.entries.append((fig_id, code, titulo, rubrica))
        if key:
            self.by_key[key] = fig_id
            self.by_key_code[key] = code
        return fig_id, code

    def ref(self, key: str, default: str = "—") -> str:
        fig = self.by_key.get(key)
        code = self.by_key_code.get(key)
        if fig and code:
            return f"{fig} ({code})"
        if fig:
            return fig
        return default


REGISTRY = FigureRegistry()

NIN_FIGURES = [
    {
        "file": "NIN-01 API Created.png",
        "key": "gw_invoke_url",
        "titulo": "HTTP API Gateway — API creada e Invoke URL",
        "texto": "HTTP API MesaTech EV1 en us-east-1. Invoke URL base del frontend. EP1 §12.",
        "rubrica": "§12 HTTP API",
        "cap": 3,
    },
    {
        "files": ["NIN-02 Rutas A.png", "NIN-02 Rutas B.png"],
        "key": "gw_rutas",
        "titulo": "Rutas expuestas en el Gateway",
        "texto": "Paths alineados al BFF y a Axios. docs/equipo/rutas-api-gateway-bff.md.",
        "rubrica": "§12 Rutas",
        "cap": 3,
    },
    {
        "files": ["NIN-03 Integrations A.png", "NIN-03 Integrations B.png"],
        "key": "gw_integracion_bff",
        "titulo": "Integración HTTP hacia el BFF en EC2",
        "texto": "Integración HTTP URI al BFF (:8080). §12.",
        "rubrica": "§12 Integración BFF",
        "cap": 3,
    },
    {
        "file": "NIN-04 JWT.png",
        "key": "gw_jwt_authorizer",
        "titulo": "JWT Authorizer (Microsoft Entra ID)",
        "texto": "Issuer Entra v2, Audience API. Sin JWT → 401. §12 / §14.",
        "rubrica": "§12 JWT Authorizer",
        "cap": 3,
    },
    {
        "file": "NIN-05 CORS.png",
        "key": "gw_cors",
        "titulo": "CORS en HTTP API",
        "texto": "Orígenes, headers y métodos para React. §12.",
        "rubrica": "§12 CORS Gateway",
        "cap": 3,
    },
    {
        "file": "NIN-06-07 401 Unauthorized.png",
        "key": "prueba_gw_sin_jwt",
        "titulo": "Gateway rechaza petición sin JWT (401)",
        "texto": "GET sin Authorization → 401. Prueba §11 #2.",
        "rubrica": "§11 #2",
        "cap": 3,
    },
    {
        "file": "NIN-06-07 401 Unauthorized.png",
        "key": "prueba_jwt_invalido",
        "titulo": "Gateway rechaza JWT inválido (401)",
        "texto": "Bearer inválido → 401. Prueba §11 #3.",
        "rubrica": "§11 #3",
        "cap": 3,
    },
    {
        "file": "NIN-08-09 Browser Network.png",
        "key": "prueba_jwt_valido",
        "titulo": "Acceso autorizado vía Gateway (200)",
        "texto": "Network: XHR a execute-api con 200. §11 #4.",
        "rubrica": "§11 #4",
        "cap": 3,
    },
    {
        "file": "NIN-08-09 Browser Network.png",
        "key": "prueba_cors_navegador",
        "titulo": "React sin error CORS hacia Gateway",
        "texto": "Sin blocked by CORS. §11 #6.",
        "rubrica": "§11 #6",
        "cap": 3,
    },
    {
        "file": "NIN-CRUD.png",
        "key": "catalogo_crud_gw",
        "titulo": "Catálogo — CRUD administrador vía Gateway",
        "texto": "POST/PUT/DELETE catálogo vía Gateway. §4.1.",
        "rubrica": "§4.1 Catálogo",
        "cap": 3,
    },
    {
        "file": "NIN-10 v1.png",
        "key": "versionamiento_v1",
        "titulo": "Versionamiento — respuesta v1",
        "texto": "GET /v1/solicitudes/mias → array JSON. §15.",
        "rubrica": "§15 v1",
        "cap": 8,
    },
    {
        "file": "NIN-11 v2.png",
        "key": "versionamiento_v2",
        "titulo": "Versionamiento — respuesta v2",
        "texto": "GET /v2/solicitudes/mias → contrato v2. §11 #8.",
        "rubrica": "§11 #8",
        "cap": 8,
    },
    {
        "file": "NIN-H1 Marca paleta logo.png",
        "alt_files": ["NIN-EXTRA.png"],
        "key": "marca_corporativa",
        "titulo": "Identidad corporativa MesaTech Cloud",
        "texto": "Paleta #1e3a5f / #0ea5e9 y logo. docs/assets/marca/guia-marca.md.",
        "rubrica": "Marca EP1",
        "cap": 9,
    },
]

NIC_FIGURES = [
    {
        "file": "NIC-03-04 Postgres y Docker.png",
        "key": "infra_postgres",
        "titulo": "PostgreSQL 16 en Docker (persistencia)",
        "texto": "EC2 DB: Postgres y docker compose. §12.",
        "rubrica": "§12 PostgreSQL",
        "cap": 4,
    },
    {
        "file": "NIC-05 BFF 8080.png",
        "key": "infra_bff",
        "titulo": "BFF desplegado — puerto 8080",
        "texto": (
            "Backend For Frontend (Spring Boot) en EC2: OAuth2 Resource Server, "
            "sin JPA ni base de datos propia; recibe tráfico del Gateway y proxifica "
            "a los microservicios. §12 / §15."
        ),
        "rubrica": "§12 BFF",
        "cap": 5,
    },
    {
        "file": "NIC-05 8081.png",
        "key": "infra_ms_solicitudes",
        "titulo": "Microservicio ms-solicitudes (:8081)",
        "texto": (
            "Dominio solicitudes: persistencia PostgreSQL, reglas de transición de "
            "estado. El BFF enruta /v1/solicitudes* hacia este MS. §12."
        ),
        "rubrica": "§12 MS solicitudes",
        "cap": 5,
    },
    {
        "file": "NIC-05 8082.png",
        "key": "infra_ms_catalogo",
        "titulo": "Microservicio ms-catalogo (:8082)",
        "texto": (
            "Catálogo categorías/prioridades; JDBC a Postgres. El BFF enruta "
            "/v1/catalogo* y aplica 403 si el rol no es administrador. §12."
        ),
        "rubrica": "§12 MS catálogo",
        "cap": 5,
    },
    {
        "file": "NIC-07 401 BFF SIN JWT.png",
        "key": "bff_401_sin_jwt",
        "titulo": "BFF responde 401 sin JWT",
        "texto": (
            "Prueba directa al BFF: petición sin header Authorization → 401. "
            "Complementa la validación en Gateway (cap. 3). §14."
        ),
        "rubrica": "§14 BFF 401",
        "cap": 5,
    },
    {
        "file": "NIC-07 202 PUBLIC.png",
        "key": "bff_ruta_publica",
        "titulo": "Ruta pública vs rutas protegidas en BFF",
        "texto": "Contraste endpoint público (p. ej. health) vs API de negocio protegida. §15.",
        "rubrica": "§15 rutas BFF",
        "cap": 5,
    },
    {
        "file": "NIC-07 - 401 Browser.png",
        "key": "bff_401_navegador",
        "titulo": "Consumo no autenticado hacia el BFF",
        "texto": "Navegador o cliente sin sesión MSAL no obtiene datos de negocio. §15.",
        "rubrica": "§15 sin auth",
        "cap": 5,
    },
    {
        "file": "NIC-08 TIMEOUT.png",
        "key": "ms_aislamiento",
        "titulo": "Microservicio no expuesto a Internet",
        "texto": "Timeout desde fuera de VPC. §15.",
        "rubrica": "§15 aislamiento MS",
        "cap": 4,
    },
    {
        "files": ["NIC-09 BFF.png", "NIC-09 MS.png", "NIC-09 DB.png"],
        "key": "infra_security_groups",
        "titulo": "Security Groups por capa",
        "texto": "Reglas BFF, MS y DB. §12.",
        "rubrica": "§12 Security Groups",
        "cap": 4,
    },
]

SKA_FIGURES = [
    {
        "file": "SKA-T01.png",
        "key": "sin_autenticacion",
        "titulo": "Acceso sin autenticación",
        "texto": "Landing con «Iniciar sesión»; sin datos de negocio. §11 #1.",
        "rubrica": "§11 #1",
        "cap": 6,
    },
    {
        "file": "SKA-UI01.png",
        "key": "ui_principal",
        "titulo": "Pantalla principal MesaTech Cloud",
        "texto": "Logo y paleta; entrada MSAL. §12.",
        "rubrica": "§12 UI",
        "cap": 6,
    },
    {
        "file": "SKA-UI02 crear solicitud.png",
        "key": "crear_solicitud",
        "titulo": "Alta de solicitud",
        "texto": "Estado inicial CREADA. §4.1.",
        "rubrica": "§4.1 crear",
        "cap": 6,
    },
    {
        "file": "SKA-UI03 mis solicitudes.png",
        "key": "listar_mias",
        "titulo": "Bandeja «Mis solicitudes»",
        "texto": "GET /v1/solicitudes/mias vía Gateway. §12.",
        "rubrica": "§12 mías",
        "cap": 6,
    },
    {
        "file": "SKA-UI04 bandeja general.png",
        "key": "listar_todas",
        "titulo": "Bandeja general (operador / admin)",
        "texto": "GET /v1/solicitudes para operador. §15.",
        "rubrica": "§15 operador",
        "cap": 6,
    },
    {
        "file": "SKA-UI05 cambio estado.png",
        "key": "cambio_estado",
        "titulo": "Actualización de estado",
        "texto": "PATCH estado y reglas MS. §4.1.",
        "rubrica": "§4.1 estados",
        "cap": 6,
    },
    {
        "file": "SKA-T05 403 cliente ver todas.png",
        "key": "auth_403_ui",
        "titulo": "Rol cliente — UI sin bandeja general",
        "texto": "Matriz de roles en navegación. §11 #5.",
        "rubrica": "§11 #5 UI",
        "cap": 7,
    },
    {
        "file": "SKA-T05b 403 Postman.png",
        "key": "auth_403_api",
        "titulo": "Rol cliente — GET /v1/solicitudes → 403",
        "texto": "BFF EntraRoles. §11 #5.",
        "rubrica": "§11 #5 API",
        "cap": 7,
    },
    {
        "file": "SKA-T03 401 Bearer invalid Postman.png",
        "key": "prueba_jwt_invalido_ska",
        "titulo": "JWT inválido rechazado (401) — Postman",
        "texto": "PATCH/GET con Bearer inválido contra API Gateway → 401. "
        "Complementa evidencia Gateway; prueba §11 #3.",
        "rubrica": "§11 #3",
        "cap": 7,
    },
    {
        "file": "SKA-T07b 409 RESUELTA sin EN_PROCESO Postman.png",
        "key": "regla_estado_resuelta",
        "titulo": "Regla de negocio — RESUELTA sin EN_PROCESO (409)",
        "texto": "PATCH de estado hacia RESUELTA desde un estado distinto de EN_PROCESO "
        "→ HTTP 409 Conflict; regla en ms-solicitudes (TransicionesEstado). §11 #7.",
        "rubrica": "§11 #7",
        "cap": 7,
    },
    {
        "file": "SKA-T07 409 transicion invalida Postman.png",
        "key": "regla_estado_cerrada",
        "titulo": "Regla de negocio — cierre inválido EN_PROCESO → CERRADA (409)",
        "texto": "PATCH hacia CERRADA sin pasar por RESUELTA → 409; la API no persiste "
        "el cambio. Misma capa de dominio que la UI. §11 #7.",
        "rubrica": "§11 #7",
        "cap": 7,
    },
    {
        "file": "SKA-T09 GET solicitudes persistencia Postman.png",
        "key": "persistencia_postman",
        "titulo": "Persistencia — GET /v1/solicitudes vía Gateway (200)",
        "texto": "Array JSON con id, título, estado y fecha desde ms-solicitudes y PostgreSQL; "
        "reconsulta independiente del frontend (Postman + Bearer). Complementa UI02/UI03. §11 #9.",
        "rubrica": "§11 #9",
        "cap": 7,
    },
]


def _paths_for(item: dict, base: Path) -> list[Path]:
    names = item.get("files") or [item["file"]]
    paths = [base / n for n in names]
    if any(p.is_file() for p in paths):
        return [p for p in paths if p.is_file()]
    for alt in item.get("alt_files", []):
        p = base / alt
        if p.is_file():
            return [p]
    return paths


def add_catalog_figure(
    doc: Document,
    item: dict,
    base: Path,
    registry: FigureRegistry | None = None,
    width_cm: float = 15.5,
) -> bool:
    reg = registry or REGISTRY
    cap = item["cap"]
    fig_id, code = reg.allocate(cap, item["titulo"], item["rubrica"], item.get("key"))
    paths = _paths_for(item, base)
    return add_figure_block(
        doc,
        fig_id,
        item["titulo"],
        item["texto"],
        item["rubrica"],
        image_paths=paths,
        width_cm=width_cm,
        evidencia_code=code,
    )


def figures_for_cap(catalog: list[dict], cap: int) -> list[dict]:
    return [f for f in catalog if f.get("cap") == cap]


def append_chapter_figures(
    doc: Document,
    base: Path,
    cap: int,
    title: str,
    intro: str,
    catalog: list[dict],
    registry: FigureRegistry | None = None,
) -> None:
    if title:
        add_chapter_heading(doc, title, intro or None)
    for item in figures_for_cap(catalog, cap):
        add_catalog_figure(doc, item, base, registry=registry)


def _para_image_blobs(paragraph) -> list[bytes]:
    blobs: list[bytes] = []
    part = paragraph.part
    for blip in paragraph._element.xpath(".//a:blip"):
        r_id = blip.get(qn("r:embed"))
        rel = part.related_parts.get(r_id)
        if rel is not None and hasattr(rel, "blob"):
            blobs.append(rel.blob)
    return blobs


def extract_entra_blocks(source: Document) -> list[dict]:
    start_markers = ("Parte A", "2. Microsoft Entra")
    stop_markers = ("Parte C", "4. Infraestructura", "Parte B", "Anexo A")
    capturing = False
    blocks: list[dict] = []
    current: dict | None = None

    for para in source.paragraphs:
        t = para.text.strip()
        if not capturing and any(t.startswith(m) for m in start_markers):
            capturing = True
            continue
        if capturing and any(t.startswith(m) for m in stop_markers):
            break
        if not capturing:
            continue

        if para.style.name in ("Heading 1", "Heading 2") and t:
            if current:
                blocks.append(current)
            current = {"titulo": t, "texto": "", "blobs": []}
            continue

        if current is None:
            continue
        if t:
            extra = current["texto"]
            current["texto"] = f"{extra} {t}".strip() if extra else t
        current["blobs"].extend(_para_image_blobs(para))

    if current:
        blocks.append(current)
    return blocks


def append_entra_from_source(
    doc: Document,
    source: Document,
    registry: FigureRegistry | None = None,
    *,
    skip_chapter_heading: bool = False,
) -> None:
    reg = registry or REGISTRY
    if not skip_chapter_heading:
        add_chapter_heading(
            doc,
            "2. Microsoft Entra ID",
            "Registro SPA + API, scope access_as_user, token v2, app roles y usuarios "
            "de demostración (contraseñas omitidas en informe).",
        )
    blocks = extract_entra_blocks(source)
    if not blocks:
        p = doc.add_paragraph()
        add_run_styled(p, "[Sin bloques Entra en la fuente.]", italic=True)
        return

    cap = 2
    for block in blocks:
        titulo = block["titulo"]
        texto = block["texto"] or "Configuración en portal Microsoft Entra ID."
        rubrica = "§12 / §14 Microsoft Entra ID"
        fig_id, code = reg.allocate(cap, titulo, rubrica)
        add_figure_block(
            doc,
            fig_id,
            titulo,
            texto,
            rubrica,
            image_blobs=block["blobs"],
            evidencia_code=code,
        )


def add_indice_evidencias(doc: Document, registry: FigureRegistry) -> None:
    p = doc.add_paragraph()
    add_run_styled(p, "Índice de evidencias", bold=True, size=13, color=COLOR_PRIMARIO)
    rows = [(fig, code, titulo[:60], rub) for fig, code, titulo, rub in registry.entries]
    add_branded_table(
        doc,
        ("Figura", "Código MT", "Descripción", "Rúbrica"),
        rows,
    )
