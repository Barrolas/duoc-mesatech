"""Estilo corporativo MesaTech Cloud para informes Word (EP1)."""

from __future__ import annotations

import io
import re
from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

# docs/assets/marca/guia-marca.md
COLOR_PRIMARIO = RGBColor(0x1E, 0x3A, 0x5F)
COLOR_ACENTO = RGBColor(0x0E, 0xA5, 0xE9)
COLOR_TEXTO = RGBColor(0x0F, 0x17, 0x2A)
COLOR_FONDO_TABLA = "1E3A5F"
COLOR_FILA_ALT = "F1F5F9"

PASSWORD_RE = re.compile(r"(Contraseña|Password)\s*:\s*.+", re.I)


def redact_text(text: str) -> str:
    return PASSWORD_RE.sub(r"\1: [omitida en informe]", text)


def setup_document(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = COLOR_TEXTO


def _shade_cell(cell, fill_hex: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill_hex)
    shd.set(qn("w:val"), "clear")
    tc_pr.append(shd)


def add_run_styled(paragraph, text: str, *, bold=False, italic=False, size=11, color=None):
    run = paragraph.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = color
    return run


def add_chapter_heading(doc: Document, title: str, intro: str | None = None) -> None:
    doc.add_page_break()
    p = doc.add_paragraph()
    add_run_styled(p, title, bold=True, size=16, color=COLOR_PRIMARIO)
    if intro:
        pi = doc.add_paragraph()
        add_run_styled(pi, intro, size=11, color=COLOR_TEXTO)
    doc.add_paragraph()


def add_subsection_heading(doc: Document, title: str) -> None:
    p = doc.add_paragraph()
    add_run_styled(p, title, bold=True, size=13, color=COLOR_PRIMARIO)
    doc.add_paragraph()


def add_body_paragraphs(doc: Document, paragraphs: Iterable[str]) -> None:
    for text in paragraphs:
        if not text.strip():
            continue
        p = doc.add_paragraph()
        add_run_styled(p, redact_text(text.strip()), size=11, color=COLOR_TEXTO)
    doc.add_paragraph()


def add_branded_table(doc: Document, headers: tuple[str, ...], rows: Iterable[tuple[str, ...]]) -> None:
    rows = [headers, *list(rows)]
    table = doc.add_table(rows=len(rows), cols=len(headers))
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell = table.rows[i].cells[j]
            cell.text = ""
            p = cell.paragraphs[0]
            if i == 0:
                _shade_cell(cell, COLOR_FONDO_TABLA)
                add_run_styled(p, val, bold=True, size=10, color=RGBColor(0xFF, 0xFF, 0xFF))
            else:
                if i % 2 == 0:
                    _shade_cell(cell, COLOR_FILA_ALT)
                add_run_styled(p, val, size=10, color=COLOR_TEXTO)
    doc.add_paragraph()


def add_figure_block(
    doc: Document,
    fig_id: str,
    titulo: str,
    texto: str,
    rubrica: str,
    image_paths: list[Path] | None = None,
    image_blobs: list[bytes] | None = None,
    width_cm: float = 15.5,
    evidencia_code: str | None = None,
) -> bool:
    """Bloque unificado: ID, título, párrafo, imagen(es), pie de figura."""
    doc.add_paragraph()
    p_id = doc.add_paragraph()
    add_run_styled(p_id, fig_id, bold=True, size=13, color=COLOR_PRIMARIO)

    p_title = doc.add_paragraph()
    add_run_styled(p_title, titulo, bold=True, size=12, color=COLOR_TEXTO)

    p_body = doc.add_paragraph()
    add_run_styled(p_body, redact_text(texto), size=11, color=COLOR_TEXTO)

    inserted = False
    paths = image_paths or []
    blobs = image_blobs or []

    panel = 0
    for path in paths:
        if not path.is_file():
            continue
        inserted = True
        if len(paths) > 1:
            panel += 1
            sub = doc.add_paragraph()
            add_run_styled(sub, f"Lámina {chr(64 + panel)}", italic=True, size=9, color=COLOR_ACENTO)
        para = doc.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para.add_run().add_picture(str(path), width=Cm(width_cm))

    for idx, blob in enumerate(blobs):
        if not blob:
            continue
        inserted = True
        para = doc.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para.add_run().add_picture(io.BytesIO(blob), width=Cm(width_cm))

    if inserted:
        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        code_part = f" · Evidencia {evidencia_code}" if evidencia_code else ""
        add_run_styled(
            cap,
            f"Pie de figura — {fig_id}{code_part}: {titulo}. {rubrica}.",
            italic=True,
            size=9,
            color=COLOR_ACENTO,
        )
    else:
        pend = doc.add_paragraph()
        add_run_styled(
            pend,
            "[Captura pendiente para esta figura.]",
            italic=True,
            size=10,
            color=COLOR_TEXTO,
        )

    doc.add_paragraph()
    return inserted


def add_portada(doc: Document) -> None:
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_run_styled(t, "MesaTech Cloud", bold=True, size=28, color=COLOR_PRIMARIO)

    s = doc.add_paragraph()
    s.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_run_styled(s, "Informe evaluación parcial 1 (EP1)", bold=True, size=16, color=COLOR_ACENTO)

    for line in (
        "Cloud Native — Duoc UC",
        "Equipo: Ari · Ninna · Nico · Skarlet",
        "Identidad, Gateway, EC2, BFF/MS, React y pruebas de autorización",
    ):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_run_styled(p, line, size=11, color=COLOR_TEXTO)

    doc.add_paragraph()
    p = doc.add_paragraph()
    run = p.add_run()
    run.add_break(WD_BREAK.PAGE)


def add_indice(doc: Document) -> None:
    p = doc.add_paragraph()
    add_run_styled(p, "Índice", bold=True, size=16, color=COLOR_PRIMARIO)
    for line in (
        "1. Contexto del caso y propuesta de solución",
        "2. Microsoft Entra ID e identidad del servicio",
        "3. API Gateway (AWS HTTP API)",
        "4. Infraestructura EC2 y PostgreSQL",
        "5. Backend BFF y microservicios",
        "6. Aplicación React y MSAL",
        "7. Autorización, roles y pruebas de API",
        "8. Versionamiento v1 / v2",
        "9. Identidad corporativa",
        "10. Flujo E2E, pruebas §11 y checklist §15",
        "11. Conclusiones",
        "Anexo A — Entorno local",
    ):
        ip = doc.add_paragraph()
        add_run_styled(ip, line, size=11, color=COLOR_TEXTO)
    doc.add_paragraph()
    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)
